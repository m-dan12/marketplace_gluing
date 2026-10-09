"""Excel-отчёт: те же листы, что в файлах склеек, плюс колонка «Почему»."""
from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .model import Card, Group


def hit_threshold(cards: list[Card], share: float = 0.2) -> int:
    """Хит: заказов за 30 дней не меньше, чем у верхних 20% карточек с продажами."""
    sold = sorted((c.orders30 for c in cards if c.orders30 > 0), reverse=True)
    if not sold:
        return 1
    return sold[max(0, int(len(sold) * share) - 1)]


def _sheet(wb, title: str, header: list[str], rows: list[list]):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(header, 1):
        ws.cell(1, i).font = Font(bold=True)
        width = max(len(str(h)), *(len(str(r[i - 1])) for r in rows[:300])) if rows else len(str(h))
        ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
    ws.freeze_panes = "A2"
    return ws


def write_report(path: Path, groups: list[Group], cards: list[Card], info: dict, flags: dict,
                 hard: list[str], profile: dict, today: date) -> None:
    thr = hit_threshold(cards)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    rows = []
    for g in groups:
        for c in sorted(g.cards, key=lambda c: (c.design, c.key)):
            rows.append([g.gid, c.brand, c.category, g.rule, c.nm_id, c.article, c.barcode,
                         c.stock_wb, c.stock_fbs, c.stock_roll, c.stock_total, c.orders30, c.rub30,
                         f"{c.prefix}{c.design}", c.key, "; ".join(g.notes)])
    _sheet(wb, "Склейки", ["ID склейки", "Бренд", "Категория", "Правило", "Артикул WB", "Артикул продавца", "Баркод",
                           "Остаток WB РФ", "Остаток FBS", "Остаток из рулона", "Итого остаток",
                           "Заказы 30д, шт", "Заказы 30д, руб", "Дизайн", "Ключ", "Почему"], rows)

    bal = []
    small_ids = set(flags["small"])
    for g in groups:
        cs = g.cards
        bal.append([g.gid, g.brand, g.category, g.rule, len(cs), len({c.design for c in cs}),
                    sum(c.orders30 for c in cs), sum(c.rub30 for c in cs), sum(c.stock_total for c in cs),
                    sum(1 for c in cs if c.orders30 == 0), sum(1 for c in cs if c.orders30 >= thr),
                    "да" if g.gid in small_ids else ""])
    _sheet(wb, "Баланс хит-неликвид", ["ID склейки", "Бренд", "Категория", "Правило", "SKU", "Дизайнов",
                                       "Заказы 30д, шт", "Заказы 30д, руб", "Итого наличие",
                                       "SKU без заказов", f"Хитов (≥{thr} заказов)", "Мало SKU"], bal)

    sizes = [len(g.cards) for g in groups]
    ctrl = [
        ["Дата", today.strftime("%d.%m.%Y")],
        ["Профиль", f"{profile['platform']} / {profile['cabinet']}"],
        ["Активных SKU", len(cards)],
        ["Склеек", len(groups)],
        ["Максимум SKU", max(sizes)],
        ["Нарушений жёстких правил", len(hard)],
        ["Малых склеек (меньше минимума)", len(flags["small"])],
        ["Склеек без заказов", len(flags["no_sales"])],
        ["Склеек из нескольких ключей", len(flags["multi_key"])],
        ["Рулонов найдено", len(info["rolls"])],
        ["SKU с наличием из рулонов", sum(1 for c in cards if c.stock_roll > 0)],
        ["Бренд не определён", len(info["unknown_brand"])],
    ] + [["Нарушение", v] for v in hard]
    _sheet(wb, "Контроль", ["Показатель", "Значение"], ctrl)

    _sheet(wb, "Реестр рулонов", ["Рулон", "Остаток, м"], [[k, v] for k, v in sorted(info["rolls"].items())])
    _sheet(wb, "Расчет тканей из рулонов",
           ["Артикул", "Рулон", "Остаток рулона, м", "Длина отреза, м", "Расчетное наличие из рулона",
            "Остаток WB РФ", "Остаток FBS", "Итого наличие", "Заказы 30д, шт"],
           [[c.article, c.roll_code, info["rolls"][c.roll_code], c.roll_len_m, c.stock_roll,
             c.stock_wb, c.stock_fbs, c.stock_total, c.orders30] for c in cards if c.roll_code])
    _sheet(wb, "Проверка малых склеек", ["ID склейки", "Бренд", "Категория", "Правило", "SKU", "Артикулы"],
           [[g.gid, g.brand, g.category, g.rule, len(g.cards), ", ".join(c.article for c in g.cards)]
            for g in groups if g.gid in small_ids])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
