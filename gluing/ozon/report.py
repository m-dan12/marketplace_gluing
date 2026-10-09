"""Excel-отчёт по склейкам Озона."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .model import Design, OGroup

HDR = PatternFill("solid", fgColor="D9E1F2")


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(list(r))
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), HDR
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, min(40, max(len(str(h)) + 2, 12)))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def write_report(path: Path | str, groups: list[OGroup], designs: list[Design], log: list[dict], quality: dict,
                 churn: dict | None, unknown: list[str]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    ctrl = [[k, v] for k, v in quality.items()]
    if churn:
        ctrl += [["— стабильность к предыдущему запуску —", ""]] + [[k, v] for k, v in churn.items()]
    ctrl += [["артикулов с неизвестным брендом (не склеены)", len(unknown)]]
    _sheet(wb, "Контроль", ["Показатель", "Значение"], ctrl, {1: 52, 2: 24})

    order = sorted(groups, key=lambda g: (g.account, g.brand, g.typ, g.leaf, g.idx))
    gid = {g.key: f"OZ-{i:03d}" for i, g in enumerate(order, 1)}
    _sheet(wb, "Склейки", ["ID", "Кабинет", "Бренд", "Тип товара", "Ветка дерева", "Часть", "Вид", "Дизайнов", "SKU",
                           "Заказы 28д", "Ключ группы (стабильный)"],
           [[gid[g.key], g.account, g.brand, g.typ, g.leaf, g.idx, g.kind, len(g.designs),
             sum(len(d.skus) for d in g.designs), sum(d.orders for d in g.designs), g.key] for g in order],
           {5: 30, 11: 60})
    rows = []
    for g in order:
        for d in g.designs:
            rows.append([gid[g.key], g.kind, g.account, g.brand, g.typ, g.leaf, d.design, d.feature,
                         ", ".join(f"{f} ({n})" for f, n in d.features.most_common()), len(d.skus), d.orders,
                         ", ".join(s.offer_id for s in d.skus)])
    _sheet(wb, "Дизайны", ["ID склейки", "Вид", "Кабинет", "Бренд", "Тип товара", "Ветка", "Дизайн", "Признак дизайна",
                           "Признаки SKU", "SKU", "Заказы 28д", "Артикулы"], rows, {6: 28, 9: 34, 12: 60})
    _sheet(wb, "Решения на узлах", ["Узел (кабинет|бренд|тип)", "Что делили", "Признак", "Дизайнов", "Решение"],
           [[x["node"], x["level"], x["name"], x["n"], x["result"]] for x in log], {1: 50, 2: 30, 5: 40})
    _sheet(wb, "Не склеено", ["Кабинет", "Бренд", "Тип товара", "Дизайн", "SKU", "Артикулы"],
           [[g.account, g.brand, g.typ, d.design, len(d.skus), ", ".join(s.offer_id for s in d.skus)]
            for g in order if g.kind == "одиночка" for d in g.designs], {6: 60})
    cuts = [[d.account, d.design, s.offer_id, s.roll_code, s.roll_m, s.cut_m, s.stock_roll, s.stock, s.stock + s.stock_roll]
            for g in order for d in g.designs for s in d.skus if s.roll_code]
    if cuts:
        _sheet(wb, "Ткани из рулонов", ["Кабинет", "Дизайн", "Артикул", "Рулон", "Метров в рулоне (FBS)", "Длина отреза, м",
                                        "Отрезов из рулона", "Остаток Озона", "Итого наличие"], cuts, {3: 24})
    if unknown:
        _sheet(wb, "Бренд не определён", ["Артикул"], [[a] for a in unknown])
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
